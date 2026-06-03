import csv
import io
from datetime import date, datetime

from django.template.loader import render_to_string
from django.utils import timezone

from apps.reports.services.analytics_service import AnalyticsService


class ExportService:
    """Handles report export generation in various formats."""

    FORMATS = ["csv", "xlsx", "pdf"]

    @staticmethod
    def generate(report, export_format):
        """Generate report content based on type and format."""
        if export_format not in ExportService.FORMATS:
            raise ValueError(f"Unsupported format: {export_format}")

        data = ExportService._fetch_report_data(report)
        file_name = f"{report.name}_{timezone.now():%Y%m%d_%H%M%S}.{export_format}"

        if export_format == "csv":
            content = ExportService._generate_csv(report.report_type, data)
        elif export_format == "xlsx":
            content = ExportService._generate_xlsx(report.report_type, data)
        else:
            content = ExportService._generate_pdf(report.report_type, data)

        return content, file_name

    @staticmethod
    def _fetch_report_data(report):
        """Fetch appropriate data based on report type."""
        svc = AnalyticsService()
        filters = report.filters

        if report.report_type == "performance":
            return svc.employee_performance(filters)
        elif report.report_type == "productivity":
            return svc.team_productivity(
                department_id=filters.get("department_id"),
                date_from=filters.get("date_from"),
                date_to=filters.get("date_to"),
            )
        elif report.report_type == "project":
            return svc.project_progress(project_id=filters.get("project_id"))
        elif report.report_type == "task":
            return svc.task_analytics(filters=filters)
        elif report.report_type == "attendance":
            return svc.attendance_analytics(
                department_id=filters.get("department_id"),
                date_from=filters.get("date_from"),
                date_to=filters.get("date_to"),
            )
        elif report.report_type == "workflow":
            return svc.workflow_efficiency(
                date_from=filters.get("date_from"),
                date_to=filters.get("date_to"),
            )
        return {}

    @staticmethod
    def _generate_csv(report_type, data):
        """Generate CSV export."""
        output = io.StringIO()
        writer = csv.writer(output)

        if report_type == "performance":
            writer.writerow(["Name", "Department", "Designation", "Total Tasks", "Completed", "Overdue", "Completion Rate (%)", "Active"])
            for row in data:
                writer.writerow([
                    row["name"], row["department"], row["designation"],
                    row["total_tasks"], row["completed_tasks"], row["overdue_tasks"],
                    row["completion_rate"], "Yes" if row["is_active"] else "No",
                ])
        elif report_type == "productivity":
            writer.writerow(["Department", "Employee Count", "Total Tasks", "Completed", "Overdue", "Completion Rate (%)"])
            for row in data:
                writer.writerow([
                    row["department"], row["employee_count"],
                    row["total_tasks"], row["completed_tasks"],
                    row["overdue_tasks"], row["completion_rate"],
                ])
        elif report_type == "project":
            writer.writerow(["Name", "Status", "Total Tasks", "Completed", "Overdue", "Progress (%)"])
            for row in data:
                writer.writerow([
                    row["name"], row["status"], row["task_count"],
                    row["completed_tasks"], row["overdue_tasks"], row["progress"],
                ])
        elif report_type == "task":
            if isinstance(data, dict):
                writer.writerow(["Metric", "Value"])
                for key, val in data.items():
                    writer.writerow([key.replace("_", " ").title(), val])
            elif isinstance(data, list):
                writer.writerow(["Name", "Status", "Priority", "Due Date"])
                for row in data:
                    writer.writerow([
                        row.get("name", ""), row.get("status", ""),
                        row.get("priority", ""), row.get("due_date", ""),
                    ])
        elif report_type == "attendance":
            if isinstance(data, dict):
                writer.writerow(["Metric", "Value"])
                for key, val in data.items():
                    writer.writerow([key.replace("_", " ").title(), val])

        return output.getvalue().encode("utf-8")

    @staticmethod
    def _generate_xlsx(report_type, data):
        """Generate Excel (XLSX) export."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = report_type.title()

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="E8B84B", end_color="E8B84B", fill_type="solid")

        if report_type == "performance":
            headers = ["Name", "Department", "Designation", "Total Tasks", "Completed", "Overdue", "Completion Rate (%)", "Active"]
            ws.append(headers)
            for row in data:
                ws.append([row["name"], row["department"], row["designation"],
                          row["total_tasks"], row["completed_tasks"], row["overdue_tasks"],
                          row["completion_rate"], "Yes" if row["is_active"] else "No"])
        elif report_type == "productivity":
            headers = ["Department", "Employee Count", "Total Tasks", "Completed", "Overdue", "Completion Rate (%)"]
            ws.append(headers)
            for row in data:
                ws.append([row["department"], row["employee_count"], row["total_tasks"],
                          row["completed_tasks"], row["overdue_tasks"], row["completion_rate"]])
        elif report_type == "project":
            headers = ["Name", "Status", "Total Tasks", "Completed", "Overdue", "Progress (%)"]
            ws.append(headers)
            for row in data:
                ws.append([row["name"], row["status"], row["task_count"],
                          row["completed_tasks"], row["overdue_tasks"], row["progress"]])
        elif report_type == "task" and isinstance(data, dict):
            ws.append(["Metric", "Value"])
            for key, val in data.items():
                ws.append([key.replace("_", " ").title(), val])
        elif report_type == "attendance" and isinstance(data, dict):
            ws.append(["Metric", "Value"])
            for key, val in data.items():
                ws.append([key.replace("_", " ").title(), val])

        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill

        for col in ws.columns:
            max_len = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = max_len + 4

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @staticmethod
    def _generate_pdf(report_type, data):
        """Generate PDF export using HTML template rendering."""
        html = render_to_string("reports/export_pdf.html", {
            "report_type": report_type,
            "data": data,
            "generated_at": timezone.now(),
        })
        try:
            from weasyprint import HTML
            pdf = HTML(string=html).write_pdf()
            return pdf
        except ImportError:
            try:
                import subprocess
                import tempfile
                with tempfile.NamedTemporaryFile(suffix=".html", delete=False, mode="w") as f:
                    f.write(html)
                    html_path = f.name
                pdf_path = html_path.replace(".html", ".pdf")
                subprocess.run(["wkhtmltopdf", html_path, pdf_path], capture_output=True)
                with open(pdf_path, "rb") as f:
                    pdf = f.read()
                import os
                os.unlink(html_path)
                os.unlink(pdf_path)
                return pdf
            except Exception:
                return html.encode("utf-8")
